# Автоматическое сохранение результатов классификации в таблицу
import docx

def autoTable(classificationReport__, path='output/classificationReport.docx'):
    """Автосохранение classificationReport в таблицу"""
    classificationReport__ = classificationReport__.split()
    classificationReport__.insert(0, '')

    i=5
    while i < len(classificationReport__)-1:
        try:
            float(classificationReport__[i])
        except:
            try:
                float(classificationReport__[i+1])
            except:
                classificationReport__[i] = classificationReport__[i] + ' ' + classificationReport__[i+1]
                classificationReport__.pop(i+1)
                continue
        i+=1
        
    classificationReport__.insert(classificationReport__.index('accuracy')+1, '')
    classificationReport__.insert(classificationReport__.index('accuracy')+1, '')

    doc = docx.Document()

    rows = len(classificationReport__)//5

    table = doc.add_table(rows=rows, cols=5)
    table.style = 'Table Grid'

    k = 0
    for i in range(0, rows):
        for j in range(0, 5):
            cell = table.cell(i, j)
            cell.text = classificationReport__[k]
            k+=1
    doc.save(path)